"""Create a fully editable Word scientific report from validated project outputs."""
from pathlib import Path
import pandas as pd,numpy as np
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
R=Path(__file__).resolve().parents[1]
OUT=R/'REPORT_Updated_Research.docx'
S=pd.read_csv(R/'research/metrics/repeated_seed_summary.csv');BIO=pd.read_csv(R/'research/metrics/biological_candidate_associations.csv');C=pd.read_csv(R/'research/metrics/lineage_adjusted_mutation_associations.csv')
orig=pd.read_csv(R/'metrics/outer_test_metrics.csv')
D=Document();sec=D.sections[0];sec.top_margin=Inches(.83);sec.bottom_margin=Inches(.80);sec.left_margin=Inches(.91);sec.right_margin=Inches(.91)
styles=D.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10);styles['Normal'].paragraph_format.space_after=Pt(6)
for name,size in [('Title',24),('Heading 1',16),('Heading 2',12.5)]:
 st=styles[name];st.font.name='Aptos Display' if name=='Title' else 'Aptos';st.font.size=Pt(size);st.font.color.rgb=RGBColor(19,48,76);st.font.bold=True
head=sec.header.paragraphs[0];head.text='CHALLENGE 27  |  PHARMACOGENOMICS RESEARCH';head.style='Caption'
foot=sec.footer.paragraphs[0];foot.text='Educational in-vitro model • Not for clinical treatment decisions';foot.style='Caption'

def para(text='',bold=False):
 p=D.add_paragraph();r=p.add_run(text);r.bold=bold;return p

def H(text,level=1):D.add_heading(text,level)
def table(headers,rows):
 t=D.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,h in enumerate(headers):t.rows[0].cells[j].text=str(h)
 for row in rows:
  cells=t.add_row().cells
  for j,v in enumerate(row):cells[j].text=str(v)
 D.add_paragraph('')

def img(name,caption,width=6.3):
 path=R/'research/figures'/name
 if not path.exists(): raise FileNotFoundError(path)
 p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(width))
 cap=D.add_paragraph(caption);cap.style='Caption'

def metric(mode,variant):
 x=S[(S['mode']==mode)&(S.variant==variant)].iloc[0];return f'{x.rmse_mean:.3f} ± {x.rmse_sd:.3f}'

p=D.add_paragraph(style='Title');p.add_run('Gene-Expression-Based Drug Sensitivity Prediction')
p=para('Challenge 27 | Expanded evaluation of jointly unseen cell lines and drug compounds');p.style='Subtitle'
para('Project type: university mini-project, reproducible computational analysis.  Updated: 13 October 2026.')
D.add_paragraph('')
H('Abstract')
para('We integrated 200 identity-checked cancer cell lines, 20 drug groups, and 4,000 real GDSC2 fitted LN_IC50 observations from uploaded CCLE/DepMap molecular profiles. The original small Random Forest improved predictions for new cells tested with familiar drugs, but underperformed the global mean for entirely new drugs. To address the original challenge’s simultaneous-generalization research gap, we ran five fixed-seed independent re-partitions for three evaluation settings: unseen cell lines, unseen drug groups, and a strict joint holdout excluding both. A common Ridge learner supported full-feature and four ablation comparisons. The full-feature Ridge attained mean RMSE 1.781, 3.358 and 3.512, respectively. Joint holdout was not improved over the global mean (3.387) by the full-feature model. Separate mutation–response associations aligned directionally with published p53/MDM2 and PI3K studies, but are retrospective and not causal. Chemical structure and independent clinical validation remain unavailable.')
H('1. Background and motivation')
para('In vitro sensitivity studies measure how cell-line growth responds to compound exposure. IC50 is the fitted concentration associated with 50% inhibition; our continuous target LN_IC50 is its natural logarithm. Lower log-IC50 usually indicates stronger inhibition under the fitted assay conditions; it does not establish patient efficacy. Genomic features include expression (RNA transcript abundance), mutation records, and gene copy-number values. Drugs are represented by putative target and pathway annotations. Prior studies show that performance often falls as holdouts move from familiar cell–drug pairs to unseen cells, unseen compounds and jointly unseen combinations [1–3].')
H('2. Data provenance and audits')
para('The original archive contains six expected files: DepMap Public 19Q4 sample metadata, expression, mutation and copy-number CSVs; GDSC2 fitted dose responses (27 October 2023); and an original README. All six members passed local CRC and SHA-256 checks. Matching used COSMIC ID plus source name/Sanger checks; 659 multi-omics matches yielded 649 conservatively eligible cells and a final 200-cell, 20-drug coverage-selected matrix. No response values were fabricated, and no LN_IC50-based inclusion rule was used. All 4,000 selected response and genomic entries required by the models are finite. Publication-source checksum authentication, complete synonym/InChIKey matching and source release concordance remain limitations.')
para('The curated matrix contains actual repeated cell–drug measurements from 200 distinct cells and 20 source-defined compound groups. Genomics and drug-response sources were released years apart; this requires cautious interpretation.')
H('3. Model inputs and leakage prevention')
para('Within each training fold, the 100 highest-variance expression features were selected using unique training cell lines. The fixed ten-gene mutation/CN panel comprises TP53, KRAS, BRAF, EGFR, PIK3CA, PTEN, APC, RB1, MYC and CDKN2A. Mutation flags mean a qualifying source variant was recorded, not definitively mutant-versus-wild-type status. Drug features tokenize the GDSC2 source PUTATIVE_TARGET and PATHWAY_NAME annotations. Numeric preprocessing and categorical fitting occur only using training entities. No response value, cell ID or drug ID enters the predictor.')
H('4. Separate and simultaneous holdout protocols')
table(['Evaluation','Final training entities','Outer held-out entities','Measured test pairs'],[
 ['New cell lines','160 cells × 20 drugs','40 previously unseen cells','800 / seed'],
 ['New drugs','200 cells × 15 drugs','5 previously unseen drug groups','1,000 / seed'],
 ['Both new','160 cells × 15 drugs','40 new cells × 5 new drugs','200 / seed']])
para('Seeds were fixed at 107, 207, 307, 407 and 507. In the both-new setting, 1,400 cross-quadrant pairs were excluded: they never became training examples that might contaminate the joint cold test. Within each seed, an entity-disjoint train-core and inner-validation partition precedes final refitting on non-test identities. Known source drug aliases are checked for cross-partition overlap, although molecular identity verification remains incomplete. Five seeds resplit the same 20 compounds and 200 cells; they are not five external cohorts.')
H('5. Original model benchmark')
table(['Original outer evaluation','Baseline RMSE','Random Forest RMSE','Random Forest MAE'],[
 ['Unseen cell lines','3.027','1.493','1.139'],
 ['Unseen drug groups','3.279','3.634','2.731']])
para('These original fixed-seed results were rechecked against saved predictions and held-out entities. A drug-specific mean baseline on the original unseen-cell task gave RMSE 1.639, reducing the Random Forest advantage to approximately 8.9%. Pearson correlation is undefined for a constant global-mean predictor. The original Random Forest is not presented as having been retested on all new joint-cold seeds.')
H('6. Repeated, directly comparable Ridge ablation')
table(['Feature set','Unseen cells RMSE','Unseen drugs RMSE','Both new RMSE'],[
 [name,metric('cell',var),metric('drug',var),metric('both',var)] for name,var in [
 ('Global mean','global_mean'),('Ridge: all features','ridge_all'),('Ridge: expression + drug','ridge_expression_drug'),('Ridge: 10-gene panel + drug','ridge_panel_drug'),('Ridge: drug descriptors only','ridge_drug_only'),('Ridge: cell omics only','ridge_cell_only')]])
para('All ablation models use Ridge(alpha=100) and identical fold-specific feature preprocessing. With joint-cold test entities, the complete multi-omics + drug-annotation Ridge averaged RMSE 3.512 ± 0.291, worse than the global training-mean 3.387 ± 0.179. Cell omics-only Ridge averaged 3.295 ± 0.202, but this is not independent evidence that omics alone suffices for novel-drug prediction. The limited target/pathway representation may extrapolate poorly across drugs.')
img('joint_cold_repeated_rmse.png','Figure 1. Fixed-seed joint-cold RMSE for the global mean and two Ridge settings.',6.35)
img('joint_cold_ablation.png','Figure 2. Strict both-new experiment, five-seed RMSE means and standard deviations for comparable Ridge feature blocks.',6.35)
H('7. Biological interpretation and literature comparison')
para('We preselected eight plausible gene–drug pairs from a fixed ten-gene mutation panel for a retrospective case study. Two-sided Mann–Whitney tests, Benjamini–Hochberg adjustments, and 999 within-lineage label permutations were recorded. These are exploratory observations from the same screened cohort; they are not independent validations, mechanistic interventions, or established biomarkers.')
table(['Drug','Gene','Variant-record / no-record','Median LN_IC50 difference','Within-lineage p'],[
 ['Nutlin-3a (-)','TP53','139 / 61','+2.689','0.001'],
 ['Selumetinib','KRAS','35 / 165','−1.008','0.046'],
 ['Selumetinib','BRAF','21 / 179','−1.005','0.036'],
 ['Taselisib','PIK3CA','16 / 184','−1.845','0.001'],
 ['Trametinib','KRAS','35 / 165','−0.638','0.171']])
para('Nutlin-3a inhibits the MDM2–p53 interaction. Prior ovarian and AML cell-line investigations identified more pronounced Nutlin response in functional p53 backgrounds [4,5]. The direction of the recorded TP53 variant association in our cohort is consistent with that mechanism, but source mutation calls do not establish each variant’s p53 functional effect. Earlier selumetinib studies linked RAS/RAF pathway context to MEK-inhibitor response, in a tissue-dependent manner [6]. Taselisib sensitivity has been associated with PIK3CA alterations in breast cancer preclinical screens [7]. Those studies provide hypotheses, not proof that our pooled associations are causal.')
img('genomic_literature_associations.png','Figure 3. Observational LN_IC50 differences by recorded variant status. No-record is not genetically confirmed wild type.',6.4)
H('8. Uncertainty and applicability domain')
para('On 200 cell-by-seed groups in the joint-cold tests, the rank correlation between mean absolute prediction error and nearest-training-cell distance in standardized selected expression space was Spearman rho 0.225 (nominal p ≈ 0.00135). Cell identities can repeat across seeds, so treating all 200 groups as fully independent overstates precision. No reliable clinical or chemical novelty threshold has been established. The original conditional bootstrap pertains to one fixed outer test; it does not replace uncertainty across independent studies.')
img('joint_cold_applicability.png','Figure 4. Error versus closest training expression profile in the jointly cold setting; descriptive applicability analysis.',6.35)
H('9. Research software demonstration')
para('The included prototype supports two modes: selecting an existing curated cell/drug pair for inference from a saved Random Forest, and loading a new laboratory molecular-profile CSV that contains the exact selected 100 expression columns plus 20 fixed panel values, together with externally supplied drug target/pathway annotations. It rejects nonfinite or incomplete profiles and issues an in-vitro, nonclinical warning. Available global feature importances describe model reliance; they do not establish individual causal contributions.')
H('10. Chemical structures, missing inputs and next experiments')
para('The source GDSC2 workbook lacks SMILES and InChIKeys. The project includes a 20-drug fillable structure template and an RDKit validator that checks canonical stereochemical structures, fingerprints and scaffold identifiers, but no SMILES or chemical benchmark metrics have been fabricated. To run chemically meaningful scaffold holdouts, the user must supply a traceable molecular mapping of all 20 drug groups to canonical SMILES and InChIKeys, verified using PubChem/GDSC/ChEMBL. A separate external drug-response cohort and, for biomarker claims, wet-lab functional assays remain necessary.')
H('11. Conclusions and integrity statement')
para('This research extension addresses the challenge’s central split-design gap with a jointly held-out drug/cell evaluation repeated over five seeds, measures multi-omics ablations with a common learner, reports error by biological context and molecular distance, and conducts a literature-guided retrospective case study. Full-feature drug-response predictions still do not outperform the global mean on jointly novel conditions. Source mismatch, limited compound descriptors, incomplete chemical identity verification, a small selected cohort, repeated tests on one underlying dataset and absent external experimental validation prohibit claims of new-drug robustness, clinical utility or discovered causal biomarkers.')
H('References')
references=[
 '[1] Wu J et al. DeepCCDS. Advanced Science (2025). https://doi.org/10.1002/advs.202416958',
 '[2] TransCDR. BMC Biology (2024). https://doi.org/10.1186/s12915-024-02023-8',
 '[3] DrugGene. BMC Bioinformatics (2024). https://doi.org/10.1186/s12859-024-05669-x',
 '[4] Nutlin-3a and TP53 in ovarian cell lines. PLOS ONE (2015). https://doi.org/10.1371/journal.pone.0135101',
 '[5] MDM2 antagonist Nutlin-3 and AML p53 status. https://pmc.ncbi.nlm.nih.gov/articles/PMC1895324/',
 '[6] Predictive markers of selumetinib response. Mol Cancer Ther (2010). https://pubmed.ncbi.nlm.nih.gov/20587667/',
 '[7] Taselisib and PIK3CA biomarkers. Mol Cancer Ther (2020). https://doi.org/10.1158/1535-7163.MCT-19-0284',
 '[8] DepMap Public 19Q4. https://figshare.com/articles/dataset/DepMap_19Q4_Public/11384241',
 '[9] PubChem PUG REST compound properties. https://pubchem.ncbi.nlm.nih.gov/docs/pug-rest',
 '[10] GDSC Drug Data, source workbook dated 27 October 2023. https://www.cancerrxgene.org/downloads/drug_data'
]
for ref in references:para(ref)
D.save(OUT)
print('Word report created',OUT,OUT.stat().st_size)
