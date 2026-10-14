
"""Regenerate the self-contained 200-cell × 20-drug dataset from the original ZIP.
Uses the frozen coverage-tie selection manifest shipped with the project.
Never consults response magnitudes for subset inclusion.
"""
import argparse,csv,io,re,zipfile
from pathlib import Path
import numpy as np,pandas as pd
from openpyxl import load_workbook
PANEL=['TP53','KRAS','BRAF','EGFR','PIK3CA','PTEN','APC','RB1','MYC','CDKN2A']
NONCODING={'Silent','Intron','IGR',"3'UTR","5'UTR",'RNA',"5'Flank","3'Flank"}

def normalize(s): return s.fillna('').astype(str).str.lower().str.replace('[^a-z0-9]','',regex=True)
def ids(z,name):
    with z.open(name) as f: return set(pd.read_csv(f,usecols=[0],dtype=str).iloc[:,0].dropna().str.strip())
def run(archive,project_root):
    root=Path(project_root);(root/'data').mkdir(parents=True,exist_ok=True)
    z=zipfile.ZipFile(archive)
    expected={'sample_info.csv','CCLE_expression_full.csv','CCLE_mutations.csv','CCLE_gene_cn.csv',
              'GDSC2_fitted_dose_response_27Oct23.xlsx','README'}
    assert expected.issubset(set(z.namelist()))  # Full CRC test is in provenance audit
    with z.open('sample_info.csv') as f:meta=pd.read_csv(f,dtype={'COSMIC_ID':'Int64'})
    with z.open('GDSC2_fitted_dose_response_27Oct23.xlsx') as f:
        wb=load_workbook(io.BytesIO(f.read()),read_only=True,data_only=True)
    ws=wb.worksheets[0];ws.reset_dimensions();rows=ws.iter_rows(values_only=True)
    head=list(next(rows));gdsc=pd.DataFrame.from_records(rows,columns=head)
    wb.close()
    assert len(gdsc)==242036 and len(head)==19
    gdsc['COSMIC_ID']=pd.to_numeric(gdsc['COSMIC_ID']).astype('Int64')
    gdsc['DRUG_ID']=pd.to_numeric(gdsc['DRUG_ID']).astype(int)
    gdsc['LN_IC50']=pd.to_numeric(gdsc['LN_IC50'])
    assert np.isfinite(gdsc['LN_IC50']).all()
    mid=meta.dropna(subset=['COSMIC_ID'])
    assert not mid.COSMIC_ID.duplicated().any()
    joined=gdsc.merge(mid[['DepMap_ID','stripped_cell_line_name','COSMIC_ID','Sanger_model_ID','lineage']],
                      on='COSMIC_ID',validate='many_to_one')
    bad=normalize(joined.CELL_LINE_NAME)!=normalize(joined.stripped_cell_line_name)
    s1=normalize(joined.SANGER_MODEL_ID);s2=normalize(joined.Sanger_model_ID)
    bad_sanger=s1.ne('')&s2.ne('')&s1.ne(s2)
    excluded=set(joined.loc[bad|bad_sanger,'DepMap_ID'])
    expr_ids=ids(z,'CCLE_expression_full.csv')
    cn_ids=ids(z,'CCLE_gene_cn.csv')
    mut_ids=set()
    with z.open('CCLE_mutations.csv') as f:
        for chunk in pd.read_csv(f,usecols=['DepMap_ID'],dtype=str,chunksize=180000):
            mut_ids.update(chunk.DepMap_ID.dropna())
    eligible=joined[joined.DepMap_ID.isin(expr_ids&cn_ids&mut_ids) &
                    ~joined.DepMap_ID.isin(excluded)].copy()
    eligible['drug_group']=eligible.DRUG_NAME.str.strip().str.casefold().str.replace(r'\s+',' ',regex=True)
    rank=eligible.groupby('drug_group').DepMap_ID.nunique()
    chosen=pd.read_csv(root/'data'/'drug_identity_annotations.csv')
    chosen_groups=chosen.drug_group.tolist()
    assert len(chosen_groups)==20 and len(set(chosen_groups))==20
    # Confirm frozen group set is an admissible top-20 coverage selection.
    mincov=int(rank.loc[chosen_groups].min())
    assert (rank.loc[chosen_groups]>=mincov).all()
    assert (rank.drop(index=chosen_groups)<=mincov).all()
    idmap=dict(zip(chosen.drug_group,chosen.canonical_drug_id))
    for k,id_ in idmap.items():
        counts=eligible[eligible.drug_group.eq(k)].groupby('DRUG_ID').DepMap_ID.nunique()
        assert int(counts.loc[id_])==int(counts.max())
    filtered=eligible[eligible.drug_group.isin(chosen_groups)].copy()
    filtered=filtered[filtered.DRUG_ID.eq(filtered.drug_group.map(idmap))]
    # Top 200 cells by coverage; tie-break ascending DepMap ID.
    coverage=filtered.groupby('DepMap_ID').drug_group.nunique().rename('coverage').reset_index()
    coverage=coverage.sort_values(['coverage','DepMap_ID'],ascending=[False,True],kind='stable')
    selected_ids=coverage.head(200).DepMap_ID.tolist()
    saved_ids=pd.read_csv(root/'data'/'selected_cell_ids.csv').DepMap_ID.tolist()
    assert selected_ids==saved_ids
    selected=filtered[filtered.DepMap_ID.isin(selected_ids)].copy()
    assert len(selected)==4000 and not selected.duplicated(['DepMap_ID','drug_group']).any()
    columns=['DepMap_ID','COSMIC_ID','SANGER_MODEL_ID','CELL_LINE_NAME','lineage','DRUG_ID',
             'DRUG_NAME','drug_group','NLME_CURVE_ID','LN_IC50','PUTATIVE_TARGET','PATHWAY_NAME']
    selected[columns].to_csv(root/'data'/'curated_responses.csv',index=False)
    # Read only 200 expression rows; check every value in the extracted 58,676 features.
    with z.open('CCLE_expression_full.csv') as f:
        exhead=next(f).decode('utf-8-sig').strip().split(',')
        erows={}
        for line in f:
            rawid,rawvals=line.split(b',',1); cell=rawid.decode().strip().strip('"')
            if cell not in selected_ids:continue
            row=np.fromstring(rawvals.decode('ascii').strip(),dtype=np.float32,sep=',')
            assert len(row)==len(exhead)-1 and np.isfinite(row).all() and cell not in erows
            erows[cell]=row
    arr=np.stack([erows[x] for x in selected_ids])
    np.savez_compressed(root/'data'/'matched_expression_full.npz',
                        cell_ids=np.asarray(selected_ids),values=arr,gene_columns=np.asarray(exhead[1:]))
    with z.open('CCLE_gene_cn.csv') as f:
        cnhead=next(f).decode('utf-8-sig').strip().split(',')
    ccols=[next(c for c in cnhead if c.startswith(g+' (')) for g in PANEL]
    with z.open('CCLE_gene_cn.csv') as f:
        cn=pd.read_csv(f,usecols=['Unnamed: 0']+ccols,
                       dtype={c:'float32' for c in ccols})
    cn=cn.rename(columns={'Unnamed: 0':'DepMap_ID'}).set_index('DepMap_ID').loc[selected_ids,ccols]
    assert np.isfinite(cn.to_numpy()).all()
    cn.to_csv(root/'data'/'copy_number_panel.csv')
    mut_rows=[]
    with z.open('CCLE_mutations.csv') as f:
        for chunk in pd.read_csv(f,usecols=['DepMap_ID','Hugo_Symbol','Variant_Classification','isDeleterious'],
                                 dtype='string',chunksize=180000):
            frame=chunk[chunk.DepMap_ID.isin(selected_ids)&chunk.Hugo_Symbol.isin(PANEL)]
            if not frame.empty:mut_rows.append(frame.copy())
    muts=pd.concat(mut_rows,ignore_index=True)
    muts['coding_variant']=~muts.Variant_Classification.isin(NONCODING)
    muts.to_csv(root/'data'/'selected_mutation_records.csv',index=False)
    mb=pd.DataFrame(0,index=selected_ids,columns=PANEL,dtype='int8')
    for cell,gene in muts[muts.coding_variant][['DepMap_ID','Hugo_Symbol']].drop_duplicates().itertuples(index=False,name=None):
        mb.loc[cell,gene]=1
    mb.index.name='DepMap_ID';mb.to_csv(root/'data'/'mutation_panel_binary.csv')
    print('RECURATION_PASSED:',arr.shape,cn.shape,mb.shape,selected.shape)
    return selected
if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('original_zip')
    parser.add_argument('project_folder')
    args=parser.parse_args()
    run(args.original_zip,args.project_folder)
