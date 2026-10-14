"""Full matched GDSC2/DepMap panel construction, leakage-safe drug group k-fold and cell k-fold.
No test-based tuning. This is a new evaluation; earlier 200x20 benchmarks are unchanged.
"""
import io,re,json,zipfile,hashlib,argparse,sys
from pathlib import Path
import numpy as np,pandas as pd
from openpyxl import load_workbook
from sklearn.model_selection import KFold
from sklearn.feature_extraction import DictVectorizer
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import Ridge
from sklearn.pipeline import make_pipeline
from scipy import sparse

GENES=['TP53','KRAS','BRAF','EGFR','PIK3CA','PTEN','APC','RB1','MYC','CDKN2A']
SEED=27013

def norm(x):return re.sub('[^a-z0-9]','',str(x).casefold())
def create(root,raw):
 root=Path(root);out=root/'research/full_gdsc';out.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(raw) as z:
  assert z.testzip() is None
  with z.open('sample_info.csv') as f:meta=pd.read_csv(f,dtype={'COSMIC_ID':'Int64'})
  with z.open('GDSC2_fitted_dose_response_27Oct23.xlsx') as f:
   w=load_workbook(io.BytesIO(f.read()),read_only=True,data_only=True)
   ws=w.worksheets[0];ws.reset_dimensions();it=ws.iter_rows(values_only=True);cols=list(next(it));gdsc=pd.DataFrame.from_records(it,columns=cols);w.close()
  assert gdsc.shape==(242036,19)
  gdsc['COSMIC_ID']=pd.to_numeric(gdsc.COSMIC_ID).astype('Int64')
  gdsc['LN_IC50']=pd.to_numeric(gdsc.LN_IC50,errors='raise')
  assert np.isfinite(gdsc.LN_IC50).all()
  gdsc['DRUG_ID']=pd.to_numeric(gdsc.DRUG_ID).astype(int)
  with z.open('CCLE_expression_full.csv') as f: expheaders=next(f).decode('utf-8-sig').strip().split(',')
  with z.open('CCLE_gene_cn.csv') as f: cnheaders=next(f).decode('utf-8-sig').strip().split(',')
  # Gate eligibility based on explicit source presence, not response magnitude.
  with z.open('CCLE_expression_full.csv') as f:
   f.readline();eid={line.split(b',',1)[0].decode().strip().strip('"') for line in f}
  with z.open('CCLE_gene_cn.csv') as f:
   f.readline();cid={line.split(b',',1)[0].decode().strip().strip('"') for line in f}
  mid=set()
  with z.open('CCLE_mutations.csv') as f:
   for c in pd.read_csv(f,usecols=['DepMap_ID'],dtype=str,chunksize=200000): mid.update(c.DepMap_ID.dropna())
  avail=eid&cid&mid
  mapping=meta.dropna(subset=['COSMIC_ID'])[['COSMIC_ID','DepMap_ID','stripped_cell_line_name','Sanger_model_ID','lineage']]
  assert mapping.COSMIC_ID.is_unique
  join=gdsc.merge(mapping,on='COSMIC_ID',validate='many_to_one')
  badname=join.apply(lambda r: norm(r.CELL_LINE_NAME)!=norm(r.stripped_cell_line_name),axis=1)
  left=join.SANGER_MODEL_ID.fillna('').astype(str).str.casefold()
  right=join.Sanger_model_ID.fillna('').astype(str).str.casefold()
  badid=left.ne('')&right.ne('')&left.ne(right)
  excluded=set(join.loc[badname|badid,'DepMap_ID'])
  valid=join[join.DepMap_ID.isin(avail)&~join.DepMap_ID.isin(excluded)].copy()
  valid['drug_group']=valid.DRUG_NAME.str.strip().str.casefold().str.replace(r'\s+',' ',regex=True)
  # Choose the best-covered numeric ID for identical normalized drug-name groups, not response-based.
  c=valid.groupby(['drug_group','DRUG_ID']).DepMap_ID.nunique().rename('n').reset_index()
  c=c.sort_values(['drug_group','n','DRUG_ID'],ascending=[True,False,True])
  canonical=c.drop_duplicates('drug_group').set_index('drug_group').DRUG_ID.to_dict()
  valid=valid[valid.DRUG_ID.eq(valid.drug_group.map(canonical))].copy()
  assert not valid.duplicated(['DepMap_ID','drug_group']).any()
  valid=valid.sort_values(['DepMap_ID','drug_group']).reset_index(drop=True)
  cellids=sorted(valid.DepMap_ID.unique())
  selected=set(cellids)
  exp_rows={}
  with z.open('CCLE_expression_full.csv') as f:
   f.readline()
   for ln,line in enumerate(f,2):
    rawid,parts=line.split(b',',1);name=rawid.decode().strip().strip('"')
    if name not in selected:continue
    values=np.fromstring(parts.decode('ascii').strip(),sep=',',dtype=np.float32)
    assert len(values)==len(expheaders)-1 and np.isfinite(values).all(),(name,ln,len(values))
    exp_rows[name]=values
  assert len(exp_rows)==len(cellids)
  X=np.stack([exp_rows[cl] for cl in cellids])
  cnc=[next(h for h in cnheaders if h.startswith(g+' (')) for g in GENES]
  with z.open('CCLE_gene_cn.csv') as f: cn=pd.read_csv(f,usecols=['Unnamed: 0']+cnc,dtype={n:'float32' for n in cnc})
  cn=cn.set_index('Unnamed: 0').loc[cellids,cnc].to_numpy('float32')
  assert np.isfinite(cn).all()
  nonsilent={'Silent','Intron','IGR',"3'UTR","5'UTR","RNA","5'Flank","3'Flank"}
  mrows=[]
  with z.open('CCLE_mutations.csv') as f:
   for chunk in pd.read_csv(f,usecols=['DepMap_ID','Hugo_Symbol','Variant_Classification'],dtype=str,chunksize=200000):
    part=chunk[chunk.DepMap_ID.isin(selected)&chunk.Hugo_Symbol.isin(GENES)&~chunk.Variant_Classification.isin(nonsilent)]
    if len(part):mrows.append(part[['DepMap_ID','Hugo_Symbol']].drop_duplicates())
  mat=pd.DataFrame(0,index=cellids,columns=GENES,dtype=np.float32)
  for df in mrows:
   for cell,gene in df.itertuples(index=False,name=None):mat.loc[cell,gene]=1
  mut=mat.to_numpy(np.float32)
  valid[['DepMap_ID','COSMIC_ID','CELL_LINE_NAME','SANGER_MODEL_ID','lineage','DRUG_ID','DRUG_NAME','drug_group','NLME_CURVE_ID','LN_IC50','PUTATIVE_TARGET','PATHWAY_NAME','MIN_CONC','MAX_CONC']].to_csv(out/'matched_responses.csv',index=False)
  np.savez_compressed(out/'matched_omics.npz',cell_ids=np.asarray(cellids),expression=X,mut=mut,cn=cn,expression_headers=np.array(expheaders[1:]))
  pd.DataFrame({'drug_group':list(canonical.keys()),'canonical_drug_id':list(canonical.values())}).to_csv(out/'canonical_drug_ids.csv',index=False)
  audit={'raw_zip_sha256':hashlib.sha256(Path(raw).read_bytes()).hexdigest(),'source_response_rows':len(gdsc),'complete_three_omics_candidates':len(avail),'excluded_conflicting_cell_ids':len(excluded),'matched_cells':len(cellids),'matched_drug_groups':valid.drug_group.nunique(),'matched_pairs':len(valid),'expression_cols':X.shape[1],'mutation_panel':GENES,'response_finite':True,'drug_ids_grouped_by_source_normalized_name_only':True,'external_chemical_alias_verification_full_catalog':False,'not_source_total_full_GDSC':True}
  (out/'audit.json').write_text(json.dumps(audit,indent=2))
  print('CURATED',audit,flush=True)
  return valid

def tokens(r):
 dic={'pathway:'+re.sub(r'\s+','_',str(r.PATHWAY_NAME).strip().casefold()):1.0,
      'drug_identity:'+str(r.drug_group):1.0}
 for t in re.findall(r'[A-Za-z][A-Za-z0-9-]*',str(r.PUTATIVE_TARGET)):
  dic['target:'+t.casefold()]=1.0
 return dic

def rmse(y,p):return float(np.sqrt(np.mean((y-p)**2)))
def corr(y,p):return float(np.corrcoef(y,p)[0,1]) if len(y)>=3 and np.std(y)>1e-10 and np.std(p)>1e-10 else np.nan

def train_eval(root):
 root=Path(root);out=root/'research/full_gdsc'
 obs=pd.read_csv(out/'matched_responses.csv')
 with np.load(out/'matched_omics.npz',allow_pickle=False) as f:
  cellids=f['cell_ids'].astype(str).tolist();expr=f['expression'];mut=f['mut'];cn=f['cn']
 lookup={name:i for i,name in enumerate(cellids)}
 output=[];predictions=[]; perdrug=[];lineages=[]
 splits=[('drug',sorted(obs.drug_group.unique())),('cell',cellids)]
 for mode,entities in splits:
  kf=KFold(n_splits=5,shuffle=True,random_state=SEED if mode=='drug' else SEED+1)
  for fold,(tri,tei) in enumerate(kf.split(entities),1):
   trainset=set(np.array(entities)[tri]); testset=set(np.array(entities)[tei]);assert trainset.isdisjoint(testset)
   col='drug_group' if mode=='drug' else 'DepMap_ID'
   tr=obs[obs[col].isin(trainset)].copy();te=obs[obs[col].isin(testset)].copy()
   assert not set(te.NLME_CURVE_ID)&set(tr.NLME_CURVE_ID)
   cell_train_idx=np.array(sorted({lookup[x] for x in tr.DepMap_ID}))
   variance=expr[cell_train_idx].var(axis=0,dtype=np.float64)
   gene_idx=np.lexsort((np.arange(len(variance)),-variance))[:100]
   # Encode each distinct drug only once; keep one-hot features sparse.
   # Avoid a massive dense matrix of mostly zero annotation columns.
   vec=DictVectorizer(sparse=True,sort=True)
   train_drugs=tr.drop_duplicates('drug_group')
   vec.fit([tokens(row) for row in train_drugs.itertuples(index=False)])
   drug_meta=obs.drop_duplicates('drug_group').set_index('drug_group').sort_index()
   drug_names=drug_meta.index
   drug_encoded=vec.transform([tokens(r) for r in drug_meta.reset_index().itertuples(index=False)]).tocsr().astype(np.float32)
   selected_expr=expr[:,gene_idx]
   def make(frame):
    idx=np.fromiter((lookup[x] for x in frame.DepMap_ID),dtype=np.int32,count=len(frame))
    didx=drug_names.get_indexer(frame.drug_group)
    if (didx<0).any():raise ValueError('Unrecognized drug in model feature matrix')
    cell=np.column_stack((selected_expr[idx],mut[idx],cn[idx])).astype(np.float32)
    # Finite matched omics is a curation invariant; no fold-level imputation needed.
    if not np.isfinite(cell).all():raise ValueError('Nonfinite omics features in fold')
    return sparse.csr_matrix(cell),drug_encoded[didx].tocsr()
   celltr,atr=make(tr);cellte,ate=make(te)
   y=tr.LN_IC50.to_numpy(float);truth=te.LN_IC50.to_numpy(float)
   control=np.full(len(te),y.mean())
   preds={'global_mean':control}
   if mode=='cell':
    pred_means=te.drug_group.map(tr.groupby('drug_group').LN_IC50.mean()).fillna(y.mean()).to_numpy(float)
    preds['drug_mean']=pred_means
   combined_train=sparse.hstack((celltr,atr),format='csr',dtype=np.float32)
   combined_test=sparse.hstack((cellte,ate),format='csr',dtype=np.float32)
   for modelname,xtr,xte in [('cell_only',celltr,cellte),('cell_plus_annotations',combined_train,combined_test)]:
    # With an intercept, scaling sparse features without subtracting the mean
    # is equivalent in linear-model geometry to centering the dense features.
    # Small solver-roundoff differences from the historic dense fits are expected.
    pipe=make_pipeline(StandardScaler(with_mean=False),Ridge(alpha=100,solver='lsqr',tol=1e-4))
    pipe.fit(xtr,y)
    preds[modelname]=pipe.predict(xte)
   tmp=te[['NLME_CURVE_ID','DepMap_ID','drug_group','lineage','LN_IC50']].copy()
   tmp['mode']=mode;tmp['fold']=fold
   for key,pr in preds.items():
    assert np.isfinite(pr).all()
    tmp['pred_'+key]=pr
    output.append({'mode':mode,'fold':fold,'model':key,'n':len(te),'train_pairs':len(tr),'train_cells':tr.DepMap_ID.nunique(),'test_cells':te.DepMap_ID.nunique(),'train_drugs':tr.drug_group.nunique(),'test_drugs':te.drug_group.nunique(),'rmse':rmse(truth,pr),'mae':float(np.mean(np.abs(truth-pr))),'pearson_r':corr(truth,pr)})
   predictions.append(tmp)
   if mode=='cell':
    for drug,g in tmp.groupby('drug_group'):
     yy=g.LN_IC50.to_numpy(float)
     for label in ['drug_mean','cell_plus_annotations']:
      pp=g['pred_'+label].to_numpy(float)
      perdrug.append({'drug_group':drug,'fold':fold,'model':label,'n':len(g),'within_drug_r':corr(yy,pp),'rmse':rmse(yy,pp),'observed_sd':float(np.std(yy,ddof=1)) if len(yy)>1 else np.nan})
    for lineage,g in tmp.groupby('lineage'):
     if len(g)<20:continue
     for label in ['drug_mean','cell_plus_annotations']:
      pp=g['pred_'+label].to_numpy(float)
      lineages.append({'lineage':lineage,'fold':fold,'model':label,'n':len(g),'rmse':rmse(g.LN_IC50.to_numpy(float),pp),'mae':float(np.mean(np.abs(g.LN_IC50.to_numpy(float)-pp)))})
   print('FOLD_COMPLETE',mode,fold,'test',len(te),'drug_groups',te.drug_group.nunique(),'RMSE',round(output[-1]['rmse'],4),flush=True)
  # save after each mode so even interruption leaves a complete mode
  pd.DataFrame(output).to_csv(out/'fold_metrics.csv',index=False)
  pd.concat(predictions,ignore_index=True).to_csv(out/'fold_predictions.csv',index=False)
  if perdrug:pd.DataFrame(perdrug).to_csv(out/'within_drug_metrics.csv',index=False)
  if lineages:pd.DataFrame(lineages).to_csv(out/'lineage_metrics.csv',index=False)
 summary=pd.DataFrame(output).groupby(['mode','model']).agg(folds=('fold','size'),rmse_mean=('rmse','mean'),rmse_sd=('rmse','std'),mae_mean=('mae','mean'),r_mean=('pearson_r','mean')).reset_index()
 summary.to_csv(out/'summary.csv',index=False)
 print('COMPLETE_FULL',summary.to_string(index=False),flush=True)

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--root',required=True);parser.add_argument('--raw',default=None);parser.add_argument('--prepare-only',action='store_true')
 args=parser.parse_args()
 if args.raw:create(args.root,args.raw)
 if not args.prepare_only: train_eval(args.root)
