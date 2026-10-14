"""Verify real secondary metrics and independent report deliverable."""
from pathlib import Path
import unittest
import numpy as np
import pandas as pd
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
class SupplementalChecks(unittest.TestCase):
    def test_original_primary_results_unchanged(self):
        m=pd.read_csv(ROOT/'metrics'/'outer_test_metrics.csv')
        expected={('unseen_cell','training_mean'):3.027192,
                  ('unseen_cell','random_forest'):1.493399,
                  ('unseen_drug','training_mean'):3.279301,
                  ('unseen_drug','random_forest'):3.634391}
        for key,r in expected.items():
            row=m[(m.split==key[0])&(m.model==key[1])].iloc[0]
            self.assertAlmostEqual(float(row.rmse),r,places=5)
    def test_stronger_baselines_are_recomputed(self):
        p=pd.read_csv(ROOT/'metrics'/'supplemental_stronger_baselines.csv')
        for setting,field,name in [('unseen_cell','DepMap_ID','per_drug_training_mean'),
                                   ('unseen_drug','drug_group','pathway_shrunk_training_mean')]:
            o=pd.read_csv(ROOT/'data'/'curated_responses.csv')
            full_train=set(pd.read_csv(ROOT/'splits'/f'{setting}_train.csv').id.astype(str))
            fit=o[o[field].astype(str).isin(full_train)]
            test=o[~o[field].astype(str).isin(full_train)]
            assert len(fit)+len(test)==4000
            glob=fit.LN_IC50.mean()
            if setting=='unseen_cell':
                prediction=test.drug_group.map(fit.groupby('drug_group').LN_IC50.mean()).fillna(glob).to_numpy()
            else:
                gp=fit.groupby('PATHWAY_NAME').LN_IC50.agg(['mean','count'])
                smooth=((gp['mean']*gp['count']+glob*100)/(gp['count']+100)).to_dict()
                prediction=test.PATHWAY_NAME.map(smooth).fillna(glob).to_numpy()
            rmse=float(np.sqrt(np.mean((test.LN_IC50.to_numpy()-prediction)**2)))
            recorded=p[(p.split==setting)&(p.phase=='outer_test_revisit')&(p.candidate==name)].iloc[0]
            self.assertAlmostEqual(rmse,recorded.rmse,places=10)
    def test_inner_only_candidate_selection_rule(self):
        candidates=pd.read_csv(ROOT/'metrics'/'supplemental_candidate_comparison.csv')
        chosen=pd.read_csv(ROOT/'metrics'/'supplemental_validation_selected_models.csv')
        assert (candidates.phase=='inner_validation').any()
        for s in ['unseen_cell','unseen_drug']:
            v=candidates[(candidates.split==s)&(candidates.phase=='inner_validation')]
            v=v.sort_values(['rmse','candidate'])
            row=chosen[chosen.split==s].iloc[0]
            self.assertEqual(row.inner_validation_selected,v.iloc[0].candidate)
            self.assertAlmostEqual(float(row.inner_validation_rmse),float(v.iloc[0].rmse),places=10)
    def test_bootstrap_intervals_and_coverage(self):
        b=pd.read_csv(ROOT/'metrics'/'conditional_cluster_bootstrap.csv')
        self.assertEqual(len(b),2)
        self.assertTrue((b.bootstrap_repeats==2000).all())
        self.assertEqual(set(b.unique_test_clusters),{40,5})
        self.assertTrue(np.isfinite(b[['rf_minus_baseline_rmse','conditional_CI_lower','conditional_CI_upper']]).all().all())
        self.assertTrue((b.conditional_CI_lower<=b.rf_minus_baseline_rmse).all())
        self.assertTrue((b.conditional_CI_upper>=b.rf_minus_baseline_rmse).all())
    def test_separate_report_deliverables(self):
        for rel in ['deliverables/scientific_report.docx','deliverables/scientific_report.pdf',
                    'deliverables/speaker_notes.docx','model_code/build_submission_documents.py']:
            self.assertGreater((ROOT/rel).stat().st_size,300)
        self.assertTrue((ROOT/'deliverables/scientific_report.pdf').read_bytes().startswith(b'%PDF-'))
    def test_ten_slide_deck_with_integrity(self):
        ppt=ROOT/'deliverables/research_presentation.pptx'
        deck=Presentation(ppt)
        self.assertEqual(len(deck.slides),10)
        self.assertTrue((ROOT/'deliverables/research_presentation.pdf').stat().st_size>100_000)
        text=' '.join(shape.text.lower() for shape in deck.slides[-1].shapes if shape.has_text_frame)
        self.assertIn('integrity',text)
        self.assertIn('uncertainty',text)
if __name__=='__main__':unittest.main()
