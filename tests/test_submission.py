"""Submission-format, leakage and stored-output checks."""
from pathlib import Path
import unittest,json
import numpy as np,pandas as pd
import joblib
from pptx import Presentation
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
class SubmissionChecks(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.responses=pd.read_csv(ROOT/'data'/'curated_responses.csv')
 def test_required_directories_and_files(self):
  for folder in ['model_code','splits','checkpoints','metrics','deliverables','figures','tests']:
   self.assertTrue((ROOT/folder).is_dir())
  for name in ['split_generation','feature_generation','training','prediction','evaluation']:
   self.assertTrue((ROOT/'model_code'/f'{name}.py').is_file())
  self.assertTrue((ROOT/'requirements.txt').is_file())
 def test_folds_partition_observations(self):
  for setting,field,counts in [
    ('unseen_cell','DepMap_ID',(2560,640,800)),
    ('unseen_drug','drug_group',(2400,600,1000))]:
   p=ROOT/'splits'/setting
   ids={}
   curves={}
   for name,count in zip(['train','validation','test'],counts):
    x=pd.read_csv(p/f'{name}.csv')
    self.assertEqual(len(x),count)
    ids[name]=set(x[field].astype(str))
    curves[name]=set(x.NLME_CURVE_ID)
   self.assertFalse(ids['train']&ids['validation'])
   self.assertFalse(ids['train']&ids['test'])
   self.assertFalse(ids['validation']&ids['test'])
   self.assertFalse(curves['train']&curves['validation'])
   self.assertFalse(curves['train']&curves['test'])
   self.assertFalse(curves['validation']&curves['test'])
   self.assertEqual(len(curves['train']|curves['validation']|curves['test']),4000)
 def test_final_train_equals_train_plus_validation(self):
  for setting,field in [('unseen_cell','DepMap_ID'),('unseen_drug','drug_group')]:
   p=ROOT/'splits'/setting
   tr=set(pd.read_csv(p/'train_entities.csv')[field].astype(str))
   val=set(pd.read_csv(p/'validation_entities.csv')[field].astype(str))
   full=set(pd.read_csv(p/'final_train_entities.csv')[field].astype(str))
   self.assertEqual(full,tr|val)
   self.assertEqual(len(pd.read_csv(p/'final_train.csv')),
       len(pd.read_csv(p/'train.csv'))+len(pd.read_csv(p/'validation.csv')))
 def test_metrics_and_inner_validation(self):
  outer=pd.read_csv(ROOT/'metrics'/'outer_test_metrics.csv')
  inner=pd.read_csv(ROOT/'metrics'/'inner_validation_metrics.csv')
  self.assertEqual(len(outer),4)
  self.assertEqual(len(inner),4)
  self.assertEqual(set(inner.split),{'unseen_cell','unseen_drug'})
  self.assertTrue(np.isfinite(inner[['rmse','mae']].to_numpy()).all())
  for split in ['unseen_cell','unseen_drug']:
   pred=pd.read_csv(ROOT/'predictions'/f'{split}_predictions.csv')
   for name in ['random_forest','training_mean']:
    p=pred[f'{name}_predicted']
    y=pred.LN_IC50
    rmse=np.sqrt(np.mean((y-p)**2))
    saved=outer[(outer.split==split)&(outer.model==name)].iloc[0]
    self.assertAlmostEqual(rmse,saved.rmse,places=10)
 def test_four_checkpoint_pipelines(self):
  for setting in ['unseen_cell','unseen_drug']:
   for model in ['random_forest','training_mean']:
    path=ROOT/'checkpoints'/f'{setting}_{model}.joblib'
    self.assertTrue(path.exists())
    d=joblib.load(path)
    self.assertIn('pipeline',d)
    self.assertIn('expression_selected_indices',d)
    self.assertIn('drug_annotation_vectorizer',d)
    self.assertEqual(len(d['expression_selected_indices']),100)
 def test_presentation_integrity(self):
  path=ROOT/'deliverables'/'research_presentation.pptx'
  pdf=ROOT/'deliverables'/'research_presentation.pdf'
  self.assertTrue(path.exists())
  self.assertTrue(pdf.exists())
  s=Presentation(path)
  self.assertEqual(len(s.slides),10)
  self.assertTrue(all(sl.has_notes_slide for sl in s.slides))
  text=' '.join(shape.text.lower() for shape in s.slides[-1].shapes if shape.has_text_frame)
  for term in ['integrity','assumptions','limitations','uncertainty','scientific']:
   self.assertIn(term,text)
 def test_requirement_pins(self):
  for line in (ROOT/'requirements.txt').read_text().splitlines():
   if line.strip():
    self.assertIn('==',line)
 def test_research_figures_are_present(self):
  png=list((ROOT/'figures').glob('*.png'))
  self.assertEqual(len(png),12)
  for f in png:
   with Image.open(f) as img:
    self.assertGreaterEqual(img.width,800)
if __name__=='__main__':
    unittest.main(verbosity=2)
