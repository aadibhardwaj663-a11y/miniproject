"""Final deliverables and primary result contract."""
import unittest
from pathlib import Path
import fitz
import numpy as np
import pandas as pd
from docx import Document
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]

class FinalReleaseChecks(unittest.TestCase):
 def test_report_is_standalone_and_reproducible(self):
  d=Document(ROOT/'deliverables/scientific_report.docx')
  content='\n'.join(p.text for p in d.paragraphs).lower()
  for key in ['challenge and scientific background','datasets and curation','prediction pipeline','experimental design','main predictive results','biological interpretation','limitations','conclusion','references']:
   self.assertIn(key,content)
  for phrase in ['1.355','1.462','2.387','2.651','0.449']:
   self.assertTrue(phrase in content or any(phrase in c.text for table in d.tables for row in table.rows for c in row.cells))
  self.assertNotIn('improved from previous model',content)
 def test_slide_count_order_and_integrity(self):
  p=Presentation(ROOT/'deliverables/research_presentation.pptx')
  self.assertEqual(len(p.slides),10)
  captions=[' '.join(sh.text.lower() for sh in slide.shapes if sh.has_text_frame) for slide in p.slides]
  for index,token in [(0,'can gene expression'),(1,'predicted'),(2,'prepare'),(3,'pipeline'),(4,'held-out'),(5,'achieve'),(6,'correlation'),(7,'difficult'),(8,'genomic'),(9,'limitations')]:
   self.assertIn(token,captions[index])
  for token in ['integrity','assumptions','scientific','uncertainty']:
   self.assertIn(token,captions[-1])
  self.assertTrue(all(sl.has_notes_slide for sl in p.slides))
 def test_speech_short_and_aligned(self):
  d=Document(ROOT/'deliverables/speaker_notes.docx')
  content='\n'.join(p.text for p in d.paragraphs)
  self.assertEqual(content.count('Slide '),10)
  self.assertLess(len(content.split()),1000)
 def test_pdf_exports(self):
  expected=[('scientific_report.pdf',(7,11)),('research_presentation.pdf',(10,10)),('speaker_notes.pdf',(1,4))]
  for name,pages in expected:
   with fitz.open(ROOT/'deliverables'/name) as file:
    self.assertGreaterEqual(len(file),pages[0]); self.assertLessEqual(len(file),pages[1])
 def test_source_metrics_match_reported_values(self):
  a=pd.read_csv(ROOT/'research/full_gdsc/summary.csv')
  for mode,model,target in [('cell','cell_plus_annotations',1.3552710216652186),('cell','drug_mean',1.4623047212764173),('drug','cell_only',2.650559535967427),('drug','cell_plus_annotations',2.3872566309863275)]:
   z=a[(a['mode']==mode)&(a.model==model)];self.assertEqual(len(z),1)
   self.assertAlmostEqual(z.iloc[0].rmse_mean,target,delta=0.001)
  p=pd.read_csv(ROOT/'research/full_gdsc/per_drug_skill_cell_holdout.csv')
  self.assertEqual((p.delta_rmse>0).sum(),229)
  self.assertAlmostEqual(p.r.median(),.449,delta=.004)
if __name__=='__main__':unittest.main()
